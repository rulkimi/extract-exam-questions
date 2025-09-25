import requests
import re
import io
from docx import Document
from docx.shared import Inches, Cm, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT

def _add_image_with_caption(paragraph, item, image_lookup, width):
    """Finds an image URL from the lookup, downloads, and inserts it with a caption."""
    # 1. Create a unique key to find the image in the lookup dictionary
    # The key normalizes whitespace in the number to ensure a match
    normalized_number = re.sub(r'\s+', '', str(item['number']))
    lookup_key = (str(item['page']), item['type'], normalized_number)
    
    image_url = image_lookup.get(lookup_key)

    if image_url:
        try:
            # Download and insert the image
            response = requests.get(image_url)
            response.raise_for_status() # Raise an exception for bad status codes
            image_stream = io.BytesIO(response.content)
            run = paragraph.add_run()
            run.add_picture(image_stream, width=width)
        except Exception as e:
            paragraph.text = f"[{item['type'].upper()} {item['number']}]"
            print(f"Failed to load image from URL {image_url}: {e}")
    else:
        # Fallback text if the image URL isn't found
        paragraph.text = f"[{item['type'].upper()} {item['number']} (URL not found)]"

    # Add caption below image
    cell = paragraph._parent
    caption_para = cell.add_paragraph()
    caption_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if item["type"] == "diagram":
        caption_para.text = f"Rajah {item['number']}"
        caption_run = caption_para.add_run(f"\nDiagram {item['number']}")
    else:  # table
        caption_para.text = f"Jadual {item['number']}"
        caption_run = caption_para.add_run(f"\nTable {item['number']}")
    caption_run.italic = True

def add_content_to_cell(cell, content, level, image_lookup):
    """Helper function to add content to a cell."""
    # Return early if content is empty or not valid
    if not content or not isinstance(content, dict):
        return  # Exit the function if content is empty or not a dictionary

    if content["type"] == "text":
        if "malay" in content["text"]:
            # Add Malay text
            malay_text = content["text"]["malay"]
            # Directly set the text of the first paragraph
            cell.paragraphs[0].text = malay_text  # Set Malay text in the first paragraph
        
        # Add English text in italics
        if "english" in content["text"]:
            english_text = content["text"]["english"]
            run = cell.paragraphs[0].add_run(f"\n{english_text}")  # Add English text in the same paragraph
            run.italic = True  # Set the run to italic
    elif content["type"] == "row":
        if "items" in content and len(content["items"]) > 0:
            cell.text = ""
            table = cell.add_table(rows=1, cols=len(content["items"]))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for idx, item in enumerate(content["items"]):
                item_cell = table.cell(0, idx)
                if item["type"] in ["diagram", "table"]:
                    item_cell.text = ""
                    paragraph = item_cell.paragraphs[0]
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    # Use the new helper function
                    width = Inches(6.0 / len(content["items"]))
                    _add_image_with_caption(paragraph, item, image_lookup, width)
                else:
                    # Handle other item types if necessary
                    item_cell.text = "[ROW ITEM]"
    elif content["type"] in ["diagram", "table"]:
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # Use the new helper function
        _add_image_with_caption(paragraph, content, image_lookup, width=Inches(6))
    elif content['type'] == "answer_space":
        if content['format'] == "line":
            for i in range(content['lines']):
                if level == 'question':
                    if i == 0:
                        cell.paragraphs[0].text = "\n……………………………………………………………………………………………………………………………………………………………"
                    else:
                        cell.paragraphs[0].add_run("\n……………………………………………………………………………………………………………………………………………………………")
                elif level == 'sub_q':
                    if i == 0:
                        cell.paragraphs[0].text = "\n……………………………………………………………………………………………………………………………………………"
                    else:
                        cell.paragraphs[0].add_run("\n……………………………………………………………………………………………………………………………………………")
                        
        elif content['format'] == "blank-space":
            cell.text = '\n\n\n\n\n'
        elif content['format'] == "multiple-choice":
            for i in range(len(content.get("options", []))):
                malay_option = content["options"][i]["malay"]
                english_option = content["options"][i]["english"]
                if i == 0:
                    cell.paragraphs[0].text = f"   [ ] {malay_option}"
                else:
                    cell.paragraphs[0].add_run(f"\n   [ ] {malay_option}")
                
                run = cell.paragraphs[0].add_run(f"\n        {english_option}\n")
                run.italic = True

def add_marks_to_cell(cell, marks):
    """Helper function to add marks to a cell."""
    cell.paragraphs[0].text = f"[{marks} markah]"
    cell.paragraphs[0].add_run(f"\n[{marks}")
    label = cell.paragraphs[0].add_run(f" marks")
    cell.paragraphs[0].add_run(f"]")
    label.italic = True
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

def replace_newlines(data):
    if isinstance(data, dict):
        for key, value in data.items():
            if isinstance(value, str):
                data[key] = value.replace("\\n", "\n")
            else:
                replace_newlines(value)
    elif isinstance(data, list):
        for item in data:
            replace_newlines(item)
            
def generate(data, image_data=None):
    replace_newlines(data)
    buffer = io.BytesIO()
    doc = Document()
    
    image_lookup = {
        (str(img['page']), img['type'], re.sub(r'\s+', '', str(img['number']))): img['url']
        for img in image_data
    }

    # Set page margins
    section = doc.sections[0]
    section.top_margin = Inches(0.5)    # Set top margin
    section.bottom_margin = Inches(0.5)  # Set bottom margin
    section.left_margin = Inches(0.5)    # Set left margin
    section.right_margin = Inches(0.5)   # Set right margin
    
    # Fill table with nested structure
    for main_q in data["main_questions"]:
        print(f"Processing main question: {main_q['number']}")
        
        # Calculate total rows needed for the current main question
        total_rows = (len(main_q["content_flow"]) +  # Add main question content rows
                    sum(len(q.get("content_flow", [])) + (1 if "marks" in q else 0) + 
                        sum(len(sq.get("content_flow", [])) + (1 if "marks" in sq else 0) 
                            for sq in q.get("sub_questions", [])) 
                        for q in main_q["questions"]))  # Add question and sub-question content rows

        print(f"Total rows for main question {main_q['number']}: {total_rows}")

        # Create a new table for each main question
        table = doc.add_table(rows=total_rows, cols=4)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Set column widths
        for row in table.rows:
            row.cells[0].width = Inches(0.3)
            row.cells[1].width = Inches(0.5)
            row.cells[2].width = Inches(0.75)
            row.cells[3].width = Inches(6.5)

        current_row = 0

        # Handle main question content
        for index, content in enumerate(main_q["content_flow"]):
            print(f"Current row: {current_row}, Content type: {content['type']}")
            if content["type"] != "questions":
                if index == 0:
                    table.rows[current_row].cells[0].text = main_q["number"]
                main_q_content_cell = table.rows[current_row].cells[1]
                main_q_content_cell.merge(table.rows[current_row].cells[-1])
                add_content_to_cell(main_q_content_cell, content, 'main_q', image_lookup)
                current_row += 1

        # Handle questions
        for question in main_q["questions"]:
            print(f"Processing question: {question['number']}")
            if current_row >= total_rows:
                print("Error: current_row exceeds total_rows")
                break  # Prevent accessing out of range

            table.rows[current_row].cells[1].text = question["number"]
            
            
            for content in question.get("content_flow", []):
                q_content_cell = table.rows[current_row].cells[2]
                q_content_cell.merge(table.rows[current_row].cells[-1])
                add_content_to_cell(q_content_cell, content, 'question', image_lookup)
                current_row += 1
                    
            if "marks" in question and "sub_questions" not in question:
                question_marks_cell = table.rows[current_row].cells[2]
                question_marks_cell.merge(table.rows[current_row].cells[-1])
                add_marks_to_cell(question_marks_cell, question["marks"])
                current_row += 1
            
            # Handle sub-questions if they exist
            for sub_q in question.get("sub_questions", []):
                print(f"Processing sub-question: {sub_q['number']}")
                
                # Check if current_row is within the valid range
                if current_row >= total_rows:
                    print(f"Error: current_row {current_row} exceeds total_rows {total_rows}")
                    break  # Prevent accessing out of range

                # Accessing the sub-question number
                table.rows[current_row].cells[2].text = sub_q["number"]
                
                for content in sub_q.get("content_flow", []):
                    sub_q_content_cell = table.rows[current_row].cells[3]
                    
                    # Check if current_row is within the valid range before accessing cells
                    if current_row >= total_rows:
                        print(f"Error: current_row {current_row} exceeds total_rows {total_rows} before adding content")
                        break  # Prevent accessing out of range
                    
                    add_content_to_cell(sub_q_content_cell, content, 'sub_q', image_lookup)
                    current_row += 1

                # Check if current_row is within the valid range before accessing marks cell
                if current_row >= total_rows:
                    print(f"Error: current_row {current_row} exceeds total_rows {total_rows} before accessing marks cell")
                    break  # Prevent accessing out of range

                if "marks" in sub_q:
                    sub_q_marks_cell = table.rows[current_row].cells[3]
                    add_marks_to_cell(sub_q_marks_cell, sub_q['marks'])
                    current_row += 1
                
            if "marks" in question and "sub_questions" in question:
                question_marks_cell = table.rows[current_row].cells[2]
                question_marks_cell.merge(table.rows[current_row].cells[-1])
                add_marks_to_cell(question_marks_cell, question["marks"])
                current_row += 1

        # Add a page break after each table
        doc.add_page_break()

    doc.save(buffer)
    buffer.seek(0)
    return buffer
    